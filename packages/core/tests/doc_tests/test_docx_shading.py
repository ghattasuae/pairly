# SPDX-License-Identifier: AGPL-3.0-only OR LicenseRef-Proofly-Commercial
# Copyright (c) 2026 Proofly Technology
"""DOCX shading: the hidden-text check compares text with the real background."""

from __future__ import annotations

import io
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from pairly_core.documents.docx import DocxParser
from pairly_core.documents.docx_shading import Shading, _adjust, theme_colours
from pairly_core.documents.registry import ParseOptions
from pairly_core.guardrails import scan_document

OPTIONS = ParseOptions(ocr_languages="eng")
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _shd(parent: object, fill: str | None = None, **attrs: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    if fill is not None:
        shd.set(qn("w:fill"), fill)
    for key, value in attrs.items():
        shd.set(qn(f"w:{key}"), value)
    parent.append(shd)  # type: ignore[attr-defined]


def _parse(tmp_path: Path, document: object) -> object:
    buf = io.BytesIO()
    document.save(buf)  # type: ignore[attr-defined]
    path = tmp_path / "doc.docx"
    path.write_bytes(buf.getvalue())
    return DocxParser().parse(path, OPTIONS)


def _flagged(doc: object) -> list[str]:
    report = scan_document(doc)  # type: ignore[arg-type]
    return [f.evidence for f in report.flagged_findings()]


def test_white_text_in_a_dark_cell_is_visible(tmp_path: Path) -> None:
    document = docx.Document()
    left, right = document.add_table(rows=1, cols=2).rows[0].cells
    _shd(left._tc.get_or_add_tcPr(), "1F2D3D")
    run = left.paragraphs[0].add_run("Skills: Python, SQL")
    run.font.color.rgb = WHITE
    right.paragraphs[0].add_run("Backend engineer")
    doc = _parse(tmp_path, document)
    backgrounds = {b.text: b.background for b in doc.pages[0].blocks}  # type: ignore[attr-defined]
    assert backgrounds["Skills: Python, SQL"] == "#1f2d3d"
    assert backgrounds["Backend engineer"] is None
    assert _flagged(doc) == []


def test_white_on_white_is_still_hidden(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph().add_run("Hidden claim one").font.color.rgb = WHITE
    cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
    _shd(cell._tc.get_or_add_tcPr(), "FFFFFF")
    cell.paragraphs[0].add_run("Hidden claim two").font.color.rgb = WHITE
    light = document.add_table(rows=1, cols=1).rows[0].cells[0]
    _shd(light._tc.get_or_add_tcPr(), "1F2D3D")
    para = light.paragraphs[0]
    _shd(para._p.get_or_add_pPr(), "FFFFFF")  # a white paragraph inside the dark cell
    para.add_run("Hidden claim three").font.color.rgb = WHITE
    flagged = " ".join(_flagged(_parse(tmp_path, document)))
    for claim in ("one", "two", "three"):
        assert f"Hidden claim {claim}" in flagged


def test_run_shading_highlight_and_auto(tmp_path: Path) -> None:
    document = docx.Document()
    shaded = document.add_paragraph().add_run("Shaded run")
    shaded.font.color.rgb = WHITE
    _shd(shaded._r.get_or_add_rPr(), "000000")
    marked = document.add_paragraph().add_run("Highlighted run")
    marked.font.color.rgb = WHITE
    highlight = OxmlElement("w:highlight")
    highlight.set(qn("w:val"), "darkBlue")
    marked._r.get_or_add_rPr().append(highlight)
    auto = document.add_paragraph().add_run("Auto shading run")
    auto.font.color.rgb = WHITE
    _shd(auto._r.get_or_add_rPr(), "auto")
    doc = _parse(tmp_path, document)
    backgrounds = {b.text: b.background for b in doc.pages[0].blocks}  # type: ignore[attr-defined]
    assert backgrounds == {
        "Shaded run": "#000000",
        "Highlighted run": "#000080",
        "Auto shading run": None,
    }
    assert _flagged(doc) == ["Auto shading run"]


def test_theme_fill_and_table_style(tmp_path: Path) -> None:
    document = docx.Document()
    cell = document.add_table(rows=1, cols=1).rows[0].cells[0]
    _shd(cell._tc.get_or_add_tcPr(), themeFill="text1")
    cell.paragraphs[0].add_run("On a theme fill").font.color.rgb = WHITE
    styled = document.add_table(rows=1, cols=1)
    styled.style = document.styles["Table Grid"]
    tcpr = styled.style.element.find(qn("w:tcPr"))
    if tcpr is None:
        tcpr = OxmlElement("w:tcPr")
        styled.style.element.append(tcpr)
    _shd(tcpr, "202020")
    styled.rows[0].cells[0].paragraphs[0].add_run("On a styled table").font.color.rgb = WHITE
    doc = _parse(tmp_path, document)
    backgrounds = {b.text: b.background for b in doc.pages[0].blocks}  # type: ignore[attr-defined]
    assert backgrounds["On a theme fill"] == "#000000"  # text1 (dk1) of the default theme
    assert backgrounds["On a styled table"] == "#202020"
    assert _flagged(doc) == []


def test_theme_helpers() -> None:
    assert theme_colours(None) == {}
    assert _adjust("#000000", "80", None) == "#7f7f7f"
    assert _adjust("#ffffff", None, "80") == "#808080"
    assert _adjust("#123456", "zz", None) == "#123456"
    shading = Shading(None, None)
    assert shading.fill(None) is None
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "nil")
    assert shading.fill(shd) is None
    unknown = OxmlElement("w:shd")
    unknown.set(qn("w:themeFill"), "accent1")
    assert shading.fill(unknown) is None
